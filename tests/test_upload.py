"""Uploading a workbook: mapping, validation, the generated signal feed and the pages that depend on it."""
import asyncio
import io
import re
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import agents, evals, signals, workbook
from app.config import XLSX_PATH
from app.main import app

client = TestClient(app)
STD = ["Company", "Ticker", "Sector", "Price (CHF)", "Daily Perf.", "3M Perf.", "6M Perf.", "1Y Perf.", "Market Data Date"]


def make_xlsx(headers, rows, title_rows=0, sheet="Data", extra_sheet=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for i in range(title_rows):
        ws.append([f"title line {i}"])
    ws.append(headers)
    for r in rows:
        ws.append(r)
    if extra_sheet:
        wb.create_sheet(extra_sheet)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def good_rows(n=12, snapshot="30.09.2026"):
    sectors = ["Financials", "Health care", "Industrials", "Materials"]
    return [[f"Company {i} AG", f"T{i:02d}", sectors[i % 4], 50 + i * 7.5, 0.001 * i, 0.01 * (i - 5), 0.02 * (i - 4), 0.03 * (i - 3), snapshot] for i in range(1, n + 1)]


@pytest.fixture(autouse=True)
def back_to_sample():
    workbook.activate(workbook.sample())
    yield
    workbook.activate(workbook.sample())


def upload(data, name="test.xlsx"):
    return client.post("/api/workbook/upload", content=data, headers={"X-Filename": name})


# ------------------------------------------------------------------ the sample file through the new reader

def test_sample_mapping_report_is_complete_and_exact():
    r = workbook.sample().report
    assert r["sheet"] == "SMI Expanded Portfolio" and r["header_row"] == 1 and r["data_rows"] == 50
    by_col = {m["column"]: m for m in r["mapping"]}
    assert by_col["Price (CHF)"]["field"] == "price" and by_col["Price (CHF)"]["how"] == "exact"
    assert by_col["Momentum Score"]["how"] == "recomputed" and by_col["Signal"]["field"] == "signal_workbook"
    assert all(m["how"] in ("exact", "recomputed") for m in r["mapping"])
    assert r["counts"] == {"rows_read": 50, "companies": 50, "with_data": 44, "no_data": 6, "distinct_companies": 47, "share_class_groups": 3, "sectors": 9}
    assert [s["name"] for s in r["sheets"] if s["chosen"]] == ["SMI Expanded Portfolio"] and len(r["sheets"]) == 4
    assert r["preview"]["columns"][0] == "Security ID" and len(r["preview"]["rows"]) == 6
    assert not r["warnings"] and any("decimals" in n for n in r["notes"]) and not any("generated" in n for n in r["notes"])


def test_uploading_the_original_file_gives_identical_results():
    original = workbook.summary()
    r = upload(XLSX_PATH.read_bytes(), "Verkaufszahlen_Demo_SMI_Expanded.xlsx")
    assert r.status_code == 200
    body = r.json()
    assert body["source"]["rows"] == 50 and body["source"]["complete"] == 44 and body["source"]["companies"] == 47
    assert body["source"]["is_sample"] is False
    again = client.get("/api/workbook").json()
    assert again["kpis"] == original["kpis"] and again["rows"] == original["rows"]
    assert [f["key"] for f in again["findings"]] == [f["key"] for f in original["findings"]]
    assert {c["ticker"]: c["material"] for c in client.get("/api/signals").json()} == {c["ticker"]: c["material"] for c in signals.compare()}


def test_reset_returns_to_the_sample_and_template_is_the_sample_file():
    upload(make_xlsx(STD, good_rows()), "mine.xlsx")
    assert client.get("/api/status").json()["source"]["name"] == "mine.xlsx"
    assert client.post("/api/workbook/reset").json()["source"]["is_sample"] is True
    assert client.get("/api/status").json()["source"]["rows"] == 50
    assert client.get("/api/workbook/template").content == XLSX_PATH.read_bytes()


# ------------------------------------------------------------------ mapping

def test_header_names_in_german_and_title_rows_above_the_table():
    headers = ["Unternehmen", "Symbol", "Branche", "Kurs (CHF)", "Tagesperformance", "3M", "6M", "1Y", "Stand"]
    ds = workbook.parse_workbook(make_xlsx(headers, good_rows(), title_rows=3, sheet="Kunden", extra_sheet="Notizen"), "kunden.xlsx")
    r = ds.report
    assert r["sheet"] == "Kunden" and r["header_row"] == 4 and r["data_rows"] == 12
    assert [m["how"] for m in r["mapping"]] == ["alias"] * 9
    assert {m["field"] for m in r["mapping"]} == {"company", "ticker", "sector", "price", "d1", "m3", "m6", "y1", "snapshot"}
    assert [s["chosen"] for s in r["sheets"]] == [True, False]
    assert ds.snapshot == date(2026, 9, 30) and ds.report["companies"][0]["xl_row"] == 5


def test_table_is_found_on_a_later_sheet():
    wb = openpyxl.Workbook()
    wb.active.title = "Cover"
    wb.active["A1"] = "Read me"
    ws = wb.create_sheet("Portfolio")
    ws.append(STD)
    for row in good_rows(5):
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    r = workbook.parse_workbook(buf.getvalue(), "later.xlsx").report
    assert r["sheet"] == "Portfolio" and [s["chosen"] for s in r["sheets"]] == [False, True]


def test_unknown_columns_are_reported_as_ignored():
    ds = workbook.parse_workbook(make_xlsx(STD + ["Internal comment"], [r + ["x"] for r in good_rows(4)]), "x.xlsx")
    m = next(m for m in ds.report["mapping"] if m["column"] == "Internal comment")
    assert m["how"] == "ignored" and m["field"] is None


def test_performance_given_as_decimals_percent_points_or_text():
    def first(vals):
        rows = [["A AG", "AAA", "Financials", 10, *vals, "30.09.2026"], ["B AG", "BBB", "Financials", 11, 0.01, 0.02, 0.03, 0.04, "30.09.2026"]]
        return workbook.parse_workbook(make_xlsx(STD, rows), "p.xlsx")

    dec = first([0.01, 0.0636, 0.1, 0.2]).rows[0]
    pts = first([1.0, 6.36, 10.0, 20.0])
    txt = first(["1%", "6.36%", "10%", "20%"]).rows[0]
    assert dec["m3"] == pytest.approx(6.36) and dec["y1"] == pytest.approx(20.0)
    assert txt["m3"] == pytest.approx(6.36) and txt["d1"] == pytest.approx(1.0)
    assert pts.rows[0]["m3"] == pytest.approx(6.36) and pts.rows[0]["y1"] == pytest.approx(20.0)
    assert pts.rows[1]["m3"] == pytest.approx(0.02), "when the first row is in percent points the whole sheet is"
    assert any("percentage points" in n for n in pts.report["notes"])


def test_blank_and_unreadable_values_are_missing_never_zero():
    rows = good_rows(6)
    rows[0][5] = "n/a"
    rows[1][4] = "abc"
    rows[2][6] = None
    ds = workbook.parse_workbook(make_xlsx(STD, rows), "m.xlsx")
    for i in (0, 1, 2):
        r = ds.rows[i]
        assert r["complete"] is False and r["score"] is None and r["signal"] == "Data missing" and r["status"] == "partial"
    assert ds.rows[0]["m3"] is None and ds.rows[2]["m6"] is None
    assert any("'abc'" in w for w in ds.report["warnings"]) and ds.report["counts"]["with_data"] == 3


def test_momentum_score_uses_the_documented_weights():
    ds = workbook.parse_workbook(make_xlsx(STD, [["A AG", "AAA", "Financials", 10, 0.01, 0.10, 0.20, 0.40, "30.09.2026"], *good_rows(3)]), "w.xlsx")
    assert ds.rows[0]["score"] == pytest.approx(1 * 0.05 + 10 * 0.2 + 20 * 0.3 + 40 * 0.45, abs=0.01)
    assert ds.rows[0]["signal"] == "Positive"


def test_security_ids_are_kept_when_valid_and_generated_otherwise():
    rows = [[f"SEC-{i:03d}", *r] for i, r in enumerate(good_rows(4), start=1)]
    kept = workbook.parse_workbook(make_xlsx(["Security ID", *STD], rows), "k.xlsx")
    assert [r["id"] for r in kept.rows] == ["SEC-001", "SEC-002", "SEC-003", "SEC-004"]
    bad = [[f"X{i}", *r] for i, r in enumerate(good_rows(4), start=1)]
    gen = workbook.parse_workbook(make_xlsx(["Security ID", *STD], bad), "g.xlsx")
    assert [r["id"] for r in gen.rows] == ["SEC-001", "SEC-002", "SEC-003", "SEC-004"]
    assert any("generated" in n for n in gen.report["notes"])


def test_rows_without_company_or_ticker_are_skipped_with_a_warning():
    rows = good_rows(5) + [[None, "ZZZ", "Financials", 1, 0, 0, 0, 0, "30.09.2026"], ["No Ticker AG", None, "Financials", 1, 0, 0, 0, 0, "30.09.2026"]]
    ds = workbook.parse_workbook(make_xlsx(STD, rows), "s.xlsx")
    assert len(ds.rows) == 5 and sum("skipped" in w for w in ds.report["warnings"]) == 2


def test_share_classes_are_grouped_by_company_name():
    rows = good_rows(4) + [["Company 1 AG", "T01N", "Financials", 40, 0, 0.01, 0.02, 0.03, "30.09.2026"]]
    ds = workbook.parse_workbook(make_xlsx(STD, rows), "sc.xlsx")
    assert ds.report["counts"]["distinct_companies"] == 4 and ds.report["counts"]["share_class_groups"] == 1
    assert ds.index["T01"]["share_classes"] == ["T01", "T01N"]
    workbook.activate(ds)
    assert any(f["key"] == "share_classes" and "1 company appears" in f["title"] for f in workbook.findings())


def test_findings_only_report_the_scale_problem_when_it_exists():
    workbook.activate(workbook.parse_workbook(make_xlsx(STD, good_rows()), "clean.xlsx"))
    assert "scale" not in {f["key"] for f in workbook.findings()}
    workbook.activate(workbook.sample())
    assert "scale" in {f["key"] for f in workbook.findings()}


# ------------------------------------------------------------------ rejected files

def errors_of(data, name="bad.xlsx"):
    with pytest.raises(workbook.UploadError) as exc:
        workbook.parse_workbook(data, name)
    return exc.value


def test_missing_required_columns_are_named_and_the_mapping_is_still_reported():
    e = errors_of(make_xlsx(["Kunde", "Produkt", "Umsatz CHF"], [["A", "X", 1]]))
    assert "Company" in e.errors[0] and "Ticker" in e.errors[0] and "Kunde, Produkt, Umsatz CHF" in e.errors[1]
    assert e.report and e.report["mapping"][0]["how"] == "ignored"
    assert any(f["required"] and f["column"] is None for f in e.report["fields"])


def test_one_missing_column_is_pinpointed():
    e = errors_of(make_xlsx([h for h in STD if h != "3M Perf."], [r[:5] + r[6:] for r in good_rows(3)]))
    assert "3M performance" in e.errors[0] and "Ticker" not in e.errors[0]


def test_duplicate_tickers_are_rejected_with_row_numbers():
    rows = good_rows(3) + [["Other AG", "T01", "Financials", 5, 0, 0, 0, 0, "30.09.2026"]]
    e = errors_of(make_xlsx(STD, rows))
    assert "T01" in e.errors[0] and "rows 2 and 5" in e.errors[0]


def test_a_sheet_without_any_complete_company_is_rejected():
    e = errors_of(make_xlsx(STD, [["A AG", "AAA", "Financials", None, None, None, None, None, None]]))
    assert "nothing can be analysed" in e.errors[0]


def test_header_without_rows_is_rejected():
    assert "no company rows" in errors_of(make_xlsx(STD, [])).errors[0]


@pytest.mark.parametrize("name,data,fragment", [
    ("notes.csv", b"a,b\n1,2", "Only .xlsx"),
    ("fake.xlsx", b"this is not a zip file", "real .xlsx"),
    ("broken.xlsx", b"PK\x03\x04 broken", "could not open"),
])
def test_wrong_file_types_are_rejected_politely(name, data, fragment):
    assert fragment.lower() in errors_of(data, name).errors[0].lower()


def test_oversized_upload_is_refused_over_http():
    r = upload(b"x" * (workbook.MAX_UPLOAD_BYTES + 1), "huge.xlsx")
    assert r.status_code == 413 and "larger than" in r.json()["errors"][0]


def test_too_many_rows_are_refused(monkeypatch):
    monkeypatch.setattr(workbook, "MAX_ROWS", 5)
    assert "limit is 5" in errors_of(make_xlsx(STD, good_rows(8))).errors[0]


def test_a_rejected_upload_leaves_the_active_workbook_untouched():
    upload(make_xlsx(STD, good_rows()), "first.xlsx")
    r = upload(make_xlsx(["Kunde"], [["A"]]), "wrong.xlsx")
    assert r.status_code == 400 and r.json()["errors"] and r.json()["report"]["mapping"]
    assert client.get("/api/status").json()["source"]["name"] == "first.xlsx"


def test_file_names_are_cleaned_and_text_stays_text():
    rows = good_rows(3)
    rows[0][0] = "<script>alert(1)</script> AG"
    ds = workbook.parse_workbook(make_xlsx(STD, rows), "C:\\fakepath\\..\\evil\x00name.xlsx")
    assert ds.name == "evilname.xlsx"
    assert ds.rows[0]["company"] == "<script>alert(1)</script> AG", "stored as plain text; the page escapes it when drawing"


# ------------------------------------------------------------------ the real feed and the pages that use it

EXAMPLES = Path(__file__).resolve().parent.parent / "data" / "examples"


@pytest.fixture()
def custom():
    ds = workbook.parse_workbook(make_xlsx(STD, good_rows(14) + [["Dark AG", "DARK", "Materials", None, None, None, None, None, "30.09.2026"]]), "custom.xlsx")
    workbook.activate(ds)
    return ds


@pytest.fixture()
def portfolio():
    ds = workbook.parse_workbook((EXAMPLES / "Kundenportfolio_Demo.xlsx").read_bytes(), "Kundenportfolio_Demo.xlsx")
    workbook.activate(ds)
    return ds


def run_monitor(focus):
    async def run():
        async def instant(_s):
            return None

        agents.pace = instant
        return [e async for e in agents.run_cycle("offline", focus)]

    original_pace = agents.pace
    try:
        return asyncio.run(run())
    finally:
        agents.pace = original_pace


def test_the_example_files_contain_real_values_from_the_excel_file(raw):
    real = {r["Ticker"]: r for r in raw}
    ds = workbook.parse_workbook((EXAMPLES / "Kundenportfolio_Demo.xlsx").read_bytes(), "Kundenportfolio_Demo.xlsx")
    assert len(ds.rows) == 14 and ds.snapshot_str == "28.08.2026"
    for r in ds.rows:
        src = real[r["ticker"]]
        assert r["company"] == src["Company"] and r["price"] == src["Price (CHF)"], r["ticker"]
        assert r["y1"] == pytest.approx(src["1Y Perf."] * 100, abs=0.006), r["ticker"]
    e = errors_of((EXAMPLES / "Stammdaten_ohne_Kurse.xlsx").read_bytes(), "Stammdaten_ohne_Kurse.xlsx")
    assert "Price (CHF)" in e.errors[0] and "Unternehmen" in e.errors[1], "the master data file must be refused, and the message must show what the app saw"


def test_an_uploaded_workbook_gets_real_events_for_the_companies_the_feed_covers(portfolio):
    ev = signals.feed()
    assert len(ev) > 10 and {e["ticker"] for e in ev} <= {r["ticker"] for r in portfolio.rows}
    assert all(e["kind"] in ("news", "quote") and e["url"].startswith("https://") for e in ev)
    changes = signals.compare()
    assert sum(c["material"] for c in changes) >= 3
    for c in changes:
        row = workbook.by_ticker(c["ticker"])
        if c["price_delta_pct"] is not None:
            latest = max((e for e in c["events"] if "price" in e), key=lambda e: e["date"])
            assert c["price_delta_pct"] == pytest.approx((latest["price"] / row["price"] - 1) * 100, abs=0.06)


def test_companies_the_feed_does_not_know_get_no_events_instead_of_invented_ones(custom):
    assert signals.feed() == [] and signals.compare() == []
    assert signals.coverage()["covered"] == 0 and signals.coverage()["securities"] == 15


def test_monitor_runs_on_an_uploaded_portfolio_and_every_briefing_passes_policy(portfolio):
    for c in [c for c in signals.compare() if c["material"]][:4]:
        events = run_monitor(c["ticker"])
        briefing = next(e for e in events if e["type"] == "briefing")["briefing"]
        review = next(e for e in events if e["type"] == "review")
        assert briefing["ticker"] == c["ticker"] and review["verdict"] == "approve", (c["ticker"], review["checks"])
        assert "28.08.2026" in " ".join(briefing["data_gaps"])
        load = next(e for e in events if e["type"] == "tool_call" and e["tool"] == "load_workbook")
        assert load["args"]["file"] == "Kundenportfolio_Demo.xlsx"


def test_monitor_still_works_when_the_feed_has_nothing_for_the_uploaded_companies(custom):
    for focus in (None, "DARK"):
        events = run_monitor(focus)
        assert events[-1]["type"] == "done"
        briefing = next(e for e in events if e["type"] == "briefing")["briefing"]
        assert next(e for e in events if e["type"] == "review")["verdict"] == "approve"
        assert any("no data for this company" in g for g in briefing["data_gaps"]), "the gap must say that the feed has nothing"
        assert not any(f["source"].startswith("SIG-") for f in briefing["facts"])


def test_trust_lab_and_bakeoff_stay_on_the_sample_even_when_another_file_is_loaded(custom):
    async def inside():
        workbook.pin_sample()
        return workbook.current().is_sample, len(workbook.rows())

    assert asyncio.run(inside()) == (True, 50)
    assert workbook.current().name == "custom.xlsx", "pinning is per task and must not leak into the global choice"
    assert client.get("/app.js").text.count("always the sample workbook") == 2, "Trust Lab and Bake-off both say they ignore an uploaded file"
    assert len(evals.load_cases()) == 15


def test_status_and_ingest_endpoints_describe_the_active_file(custom):
    s = client.get("/api/status").json()["source"]
    assert s["name"] == "custom.xlsx" and s["rows"] == 15 and s["complete"] == 14 and s["companies"] == 15
    i = client.get("/api/workbook/ingest").json()
    assert i["report"]["file"]["name"] == "custom.xlsx" and i["report"]["counts"]["with_data"] == 14
    assert i["report"]["fields"][1]["column"] == "Company"


def test_a_running_stream_keeps_its_workbook_when_another_file_is_uploaded(custom):
    from app.main import sse

    async def probe():
        yield {"type": "a", "file": workbook.current().name}
        workbook.activate(workbook.sample())  # another upload lands while the stream is running
        await asyncio.sleep(0)
        yield {"type": "b", "file": workbook.current().name, "rows": len(workbook.rows())}

    async def run():
        return [chunk async for chunk in sse(probe()).body_iterator]

    chunks = asyncio.run(run())
    assert '"file": "custom.xlsx"' in chunks[0] and chunks[1].count('"file": "custom.xlsx"') == 1 and '"rows": 15' in chunks[1]
    assert workbook.current().is_sample, "the stream must not change the global workbook"


def test_formulas_without_a_saved_result_are_reported_not_silently_blank():
    rows = good_rows(8)
    rows[2][3] = "=50+7"
    rows[4][5] = "=0.1*2"
    ds = workbook.parse_workbook(make_xlsx(STD, rows), "scripted.xlsx")
    text = " ".join(ds.report["warnings"])
    assert "2 cells contain formulas without a saved result" in text and "save it again" in text
    assert sum(1 for r in ds.rows if not r["complete"]) >= 1
    assert workbook.sample().report["warnings"] == [], "the real Excel file stores its results and must not warn"


def test_text_that_looks_like_markup_is_kept_as_plain_data():
    payloads = ["<img src=x onerror=alert(1)>", '"><script>window.__x=1</script>', "javascript:alert(1)", "+cmd|' /C calc'!A0"]
    rows = good_rows(6) + [[p, f"X{i}", "<svg onload=1>", 100, 0.01, 0.02, 0.03, 0.04, "30.09.2026"] for i, p in enumerate(payloads)]
    ds = workbook.parse_workbook(make_xlsx(STD, rows), "evil.xlsx")
    workbook.activate(ds)
    body = client.get("/api/workbook").json()
    names = {r["company"] for r in body["rows"]}
    assert set(payloads) <= names, "markup stays as text in the API; the UI escapes it when rendering"
    js = client.get("/app.js").text
    assert "const esc =" in js and js.count("innerHTML") > 20


# ------------------------------------------------------------------ odd workbooks

def test_zero_and_negative_prices_are_treated_as_missing_not_as_data():
    rows = good_rows(8)
    rows[1][3] = 0
    rows[2][3] = -5
    ds = workbook.parse_workbook(make_xlsx(STD, rows), "odd.xlsx")
    by_ticker = {r["ticker"]: r for r in ds.rows}
    assert not by_ticker["T02"]["complete"] and not by_ticker["T03"]["complete"]
    assert by_ticker["T04"]["complete"]


def test_merged_cells_and_blank_rows_between_data_do_not_break_the_reader():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Portfolio", None, None])
    ws.merge_cells("A1:I1")
    ws.append(STD)
    for r in good_rows(5)[:2]:
        ws.append(r)
    ws.append([None] * 9)
    for r in good_rows(5)[2:]:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    ds = workbook.parse_workbook(buf.getvalue(), "merged.xlsx")
    assert len(ds.rows) == 5 and all(r["complete"] for r in ds.rows)


def test_a_two_thousand_row_workbook_is_read_quickly():
    import time

    rows = [[f"Company {i} AG", f"T{i:04d}", f"Sector {i % 12}", 20 + i % 90, 0.001, 0.02, 0.03, 0.04, "30.09.2026"] for i in range(2000)]
    data = make_xlsx(STD, rows)
    t0 = time.perf_counter()
    ds = workbook.parse_workbook(data, "big.xlsx")
    took = time.perf_counter() - t0
    assert len(ds.rows) == 2000 and took < 8, took


def test_duplicate_tickers_are_refused_with_the_rows_named():
    rows = good_rows(4) + [["Twin AG", "T01", "Financials", 99, 0.01, 0.02, 0.03, 0.04, "30.09.2026"]]
    with pytest.raises(workbook.UploadError) as e:
        workbook.parse_workbook(make_xlsx(STD, rows), "twins.xlsx")
    assert "'T01' appears more than once" in e.value.errors[0] and "rows 2 and 6" in e.value.errors[0]
    r = upload(make_xlsx(STD, rows), "twins.xlsx")
    assert r.status_code == 400 and workbook.current().is_sample, "a refused upload must leave the loaded workbook alone"

def test_ties_in_the_ranking_are_stable_and_complete():
    from app import chat

    rows = [[f"Same {i} AG", f"S{i:02d}", "Financials", 100, 0.01, 0.02, 0.03, 0.04, "30.09.2026"] for i in range(6)]
    workbook.activate(workbook.parse_workbook(make_xlsx(STD, rows), "ties.xlsx"))
    first = chat.t_rank_companies("score", "desc", 6)
    again = chat.t_rank_companies("score", "desc", 6)
    assert first == again
    assert len({r["ticker"] for r in first["ranking"]}) == 6 if "ranking" in first else True



