"""Shared fixtures. The Excel file is read here with plain openpyxl, independently of app code."""
import hashlib
from collections import defaultdict

import openpyxl
import pytest

from app.config import XLSX_PATH

MISSING_TICKERS = {"AMRZ", "GALD", "SDZ", "ACLN", "SQN", "SUNN"}


def file_hash():
    return hashlib.sha256(XLSX_PATH.read_bytes()).hexdigest()


@pytest.fixture(scope="session")
def wb_values():
    return openpyxl.load_workbook(XLSX_PATH, data_only=True)


@pytest.fixture(scope="session")
def wb_formulas():
    return openpyxl.load_workbook(XLSX_PATH, data_only=False)


def _records(ws):
    rows = list(ws.iter_rows(values_only=True))
    headers = rows[0]
    return [dict(zip(headers, r)) for r in rows[1:] if r[0]]


@pytest.fixture(scope="session")
def raw(wb_values):
    """Portfolio sheet as a list of dicts keyed by the Excel header."""
    return _records(wb_values["SMI Expanded Portfolio"])


@pytest.fixture(scope="session")
def raw_complete(raw):
    return [r for r in raw if r["Price (CHF)"] is not None]


@pytest.fixture(scope="session")
def sector_groups(raw_complete):
    groups = defaultdict(list)
    for r in raw_complete:
        groups[r["Sector"]].append(r)
    return groups


@pytest.fixture(scope="session", autouse=True)
def excel_is_never_modified():
    before = file_hash()
    yield
    assert file_hash() == before, "The Excel file was modified during the test run"
