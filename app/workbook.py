"""Workbook loading, mapping, validation and analysis.

The active dataset is the sample file until someone uploads another one. Uploaded files are parsed in
memory, never written to disk and never modified.
"""
import hashlib
import io
import math
import re
from collections import Counter, defaultdict
from contextvars import ContextVar
from datetime import date, datetime
from functools import lru_cache
from statistics import mean

import openpyxl

from .config import XLSX_PATH

POSITIVE_AT, NEGATIVE_AT = 20.0, -10.0
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
HEADER_SCAN_ROWS = 10
SAMPLE_NAME = "Verkaufszahlen_Demo_SMI_Expanded.xlsx"

# key, label, required, accepted header names (first one is the canonical name), what the app does with it
FIELDS = [
    ("id", "Security ID", False, ["security id", "id", "security"], "Row reference used for citations (SEC-001)"),
    ("company", "Company", True, ["company", "company name", "name", "unternehmen", "firma"], "Shown everywhere; share classes with the same name are grouped"),
    ("ticker", "Ticker", True, ["ticker", "symbol", "kürzel", "valor"], "Unique short code that identifies the security"),
    ("index", "Index", False, ["index"], "Index membership"),
    ("sector", "Sector", True, ["sector", "industry", "branche", "sektor"], "Sector comparison and account-team routing"),
    ("canton", "Canton", False, ["canton", "kanton", "region"], "Location context"),
    ("isin", "ISIN", False, ["isin"], "Identifier shown in the table"),
    ("price", "Price (CHF)", True, ["price (chf)", "price", "kurs (chf)", "kurs", "last price", "close"], "Latest price in the snapshot"),
    ("d1", "Daily performance", True, ["daily perf.", "daily performance", "1d perf.", "1d", "tagesperformance"], "Weight 5% in the momentum score"),
    ("m3", "3M performance", True, ["3m perf.", "3m performance", "3m", "3 months", "3-monats-performance"], "Weight 20% in the momentum score"),
    ("m6", "6M performance", True, ["6m perf.", "6m performance", "6m", "6 months", "6-monats-performance"], "Weight 30% in the momentum score"),
    ("y1", "1Y performance", True, ["1y perf.", "1y performance", "1y", "12 months", "1-jahres-performance"], "Weight 45% in the momentum score"),
    ("snapshot", "Market data date", False, ["market data date", "as of", "date", "stand", "datum"], "Date of the snapshot, used to judge how old the data is"),
    ("coverage", "Data coverage", False, ["data coverage", "coverage"], "Informational"),
    ("source_url", "Source URL", False, ["source url", "source", "url", "quelle"], "Where the numbers came from"),
    ("signal_workbook", "Signal in workbook", False, ["signal"], "Kept to compare with the app's own signal"),
]
FIELD = {f[0]: f for f in FIELDS}
RECOMPUTED_HEADERS = {"momentum score"}
PERF_KEYS = ("d1", "m3", "m6", "y1")
BLANKS = {"", "n/a", "na", "-", "–", "—", "none", "null", "#n/a"}

SECTOR_TEAMS = {
    "Health care": "Life Sciences account team",
    "Industrials": "Industrial & Manufacturing account team",
    "Financials": "Financial Services account team",
    "Materials": "Industrial & Manufacturing account team",
    "Retail staples": "Consumer & Retail account team",
    "Retail discretionary": "Consumer & Retail account team",
    "Communication services": "Telecom & Media account team",
    "Information technology": "Technology account team",
    "Real estate": "Real Estate & Infrastructure account team",
}


class UploadError(Exception):
    def __init__(self, errors, report=None):
        super().__init__("; ".join(errors))
        self.errors = errors
        self.report = report


def signal_for(score_pts):
    if score_pts is None:
        return "Data missing"
    if score_pts >= POSITIVE_AT:
        return "Positive"
    if score_pts <= NEGATIVE_AT:
        return "Negative"
    return "Mixed"


# ----------------------------------------------------------------------------- parsing helpers

def _norm(h):
    return re.sub(r"\s+", " ", str(h).strip()).casefold()


def _num(v):
    """Returns (value, readable). Blank or n/a is (None, True); unreadable text is (None, False)."""
    if v is None or (isinstance(v, str) and v.strip().lower() in BLANKS):
        return None, True
    if isinstance(v, bool):
        return None, False
    if isinstance(v, (int, float)):
        return (float(v), True) if math.isfinite(v) else (None, False)
    s = re.sub(r"[\s'’]", "", str(v))
    pct = s.endswith("%")
    s = s.rstrip("%").replace(",", ".")
    try:
        x = float(s)
    except ValueError:
        return None, False
    if not math.isfinite(x):
        return None, False
    return (x / 100 if pct else x), True


def _text(v, limit=200):
    if v is None:
        return None
    s = re.sub(r"[\x00-\x1f]", " ", str(v)).strip()
    return s[:limit] or None


def _pct(v):
    return None if v is None else round(v * 100, 2)


def _parse_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y"):
            try:
                return datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                continue
    return None


def _show(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:.4f}".rstrip("0").rstrip(".")
    if isinstance(v, (datetime, date)):
        return v.strftime("%d.%m.%Y")
    s = str(v)
    return s if len(s) <= 28 else s[:27] + "…"


def _match_headers(headers):
    """Maps each header cell to an app field. Returns {column_index: (field_key, how)}."""
    taken, out = set(), {}
    for i, h in enumerate(headers):
        if h is None or not str(h).strip():
            continue
        n = _norm(h)
        for key, _label, _req, aliases, _desc in FIELDS:
            if key in taken:
                continue
            if n in aliases:
                out[i] = (key, "exact" if n == aliases[0] else "alias")
                taken.add(key)
                break
    return out


def _find_table(wb):
    """Finds the sheet and header row that hold the portfolio. Returns (ws, header_row, headers, matches)."""
    best = None
    required = {k for k, _l, req, _a, _d in FIELDS if req}
    for ws in wb.worksheets:
        for r_no, row in enumerate(ws.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS, values_only=True), start=1):
            matches = _match_headers(row)
            found = {k for k, _how in matches.values()}
            score = len(found & required)
            if best is None or score > best[0]:
                best = (score, ws, r_no, list(row), matches)
            if required <= found:
                return ws, r_no, list(row), matches
    return best[1], best[2], best[3], best[4] if best else (wb.worksheets[0], 1, [], {})


# ----------------------------------------------------------------------------- dataset

class Dataset:
    def __init__(self, name, sha, size, is_sample, rows, snapshot, report):
        self.name, self.sha, self.size, self.is_sample = name, sha, size, is_sample
        self.rows, self.snapshot, self.report = rows, snapshot, report
        self.index = {r["ticker"].upper(): r for r in rows}
        self.feed = None  # filled lazily by app.signals

    @property
    def snapshot_str(self):
        return self.snapshot.strftime("%d.%m.%Y")

    def info(self):
        complete = sum(1 for r in self.rows if r["complete"])
        return {
            "name": self.name, "is_sample": self.is_sample, "sheet": self.report["sheet"], "rows": len(self.rows),
            "complete": complete, "companies": len({r["company"] for r in self.rows}), "snapshot": self.snapshot_str,
            "size_kb": round(self.size / 1024, 1), "warnings": self.report["warnings"],
        }


def _formulas_without_result(data, sheet, header_row, columns):
    """Counts formula cells in the mapped columns that have no stored result (files written by scripts)."""
    try:
        formulas = openpyxl.load_workbook(io.BytesIO(data), data_only=False)[sheet]
        values = openpyxl.load_workbook(io.BytesIO(data), data_only=True)[sheet]
    except Exception:  # noqa: BLE001
        return 0
    count = 0
    for row in formulas.iter_rows(min_row=header_row + 1, max_row=min(formulas.max_row, header_row + MAX_ROWS + 1)):
        for i in columns:
            cell = row[i] if i < len(row) else None
            if cell is not None and isinstance(cell.value, str) and cell.value.startswith("=") and values.cell(row=cell.row, column=cell.column).value is None:
                count += 1
    return count


def parse_workbook(data, name, is_sample=False):
    name = re.sub(r"[\x00-\x1f]", "", name.replace("\\", "/").split("/")[-1])[:80] or "workbook.xlsx"
    if len(data) > MAX_UPLOAD_BYTES:
        raise UploadError([f"The file is {len(data) / 1024 / 1024:.1f} MB. The limit is {MAX_UPLOAD_BYTES // 1024 // 1024} MB."])
    if not name.lower().endswith(".xlsx"):
        raise UploadError(["Only .xlsx files are supported. In Excel use File > Save As > Excel Workbook (.xlsx)."])
    if data[:2] != b"PK":
        raise UploadError(["This does not look like a real .xlsx file. Re-save it from Excel and try again."])
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise UploadError([f"Excel could not open this file ({type(exc).__name__})."]) from exc

    sha = hashlib.sha256(data).hexdigest()
    ws, header_row, headers, matches = _find_table(wb)
    records = [(n, r) for n, r in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1)
               if any(c is not None and str(c).strip() for c in r)]
    mapping = _mapping_report(headers, matches, records)
    base_report = {
        "file": {"name": name, "size_kb": round(len(data) / 1024, 1), "sha": sha[:10], "is_sample": is_sample},
        "sheets": [{"name": s.title, "rows": s.max_row, "cols": s.max_column, "chosen": s.title == ws.title} for s in wb.worksheets],
        "sheet": ws.title, "header_row": header_row, "data_rows": len(records), "mapping": mapping,
        "fields": _field_report(headers, matches), "preview": _preview(headers, records), "warnings": [], "notes": [],
        "derived": [
            {"label": "Momentum score", "how": "5% daily + 20% 3M + 30% 6M + 45% 1Y, in percentage points"},
            {"label": "Signal", "how": f"Positive from {POSITIVE_AT:.0f} points, Negative up to {NEGATIVE_AT:.0f}, otherwise Mixed"},
            {"label": "Data completeness", "how": "Complete only if price and all four performance values are present"},
        ],
    }

    present = {k for k, _how in matches.values()}
    missing_required = [FIELD[k][1] for k, f in FIELD.items() if f[2] and k not in present]
    if missing_required:
        found = ", ".join(str(h) for h in headers if h) or "no header row"
        raise UploadError([f"Required columns not found on sheet '{ws.title}': {', '.join(missing_required)}.",
                           f"Columns the app can see: {found}."], base_report)
    if len(records) > MAX_ROWS:
        raise UploadError([f"The sheet has {len(records)} data rows. The limit is {MAX_ROWS}."], base_report)

    col = {m[0]: i for i, m in matches.items()}
    warnings, notes, rows, bad_cells = [], [], [], 0
    perf_raw = []
    parsed = []
    for xl_row, cells in records:
        def cell(key, _cells=cells):
            i = col.get(key)
            return _cells[i] if i is not None and i < len(_cells) else None

        company, ticker = _text(cell("company")), _text(cell("ticker"))
        if not company or not ticker:
            warnings.append(f"Row {xl_row} skipped: company or ticker is empty.")
            continue
        vals = {}
        for key in ("price", *PERF_KEYS):
            v, ok = _num(cell(key))
            if ok and key == "price" and v is not None and v <= 0:
                bad_cells += 1
                if bad_cells <= 5:
                    warnings.append(f"Row {xl_row}: price {v:g} is not above zero and is treated as missing.")
                v = None
            elif not ok:
                bad_cells += 1
                if bad_cells <= 5:
                    warnings.append(f"Row {xl_row}: '{_text(cell(key), 20)}' in {FIELD[key][1]} is not a number and is treated as missing.")
            vals[key] = v
        perf_raw += [vals[k] for k in PERF_KEYS if vals[k] is not None]
        parsed.append((xl_row, company, ticker, cell, vals))
    if bad_cells > 5:
        warnings.append(f"{bad_cells - 5} more unreadable values were treated as missing.")
    lost = _formulas_without_result(data, ws.title, header_row, list(matches))
    if lost:
        warnings.insert(0, f"{lost} cell{' contains a formula' if lost == 1 else 's contain formulas'} without a saved result, so {'it is' if lost == 1 else 'they are'} read as blank. "
                           "The file was probably created by a script. Open it in Excel and save it again to store the values.")

    percent_points = bool(perf_raw) and max(abs(v) for v in perf_raw) > 5
    if percent_points:
        notes.append("Performance values look like percentages (for example 6.36), so they were read as percentage points.")
    else:
        notes.append("Performance values are decimals (for example 0.0636) and are shown as percentages (6.36%).")

    ids = [_text(c("id"), 30) for _r, _co, _t, c, _v in parsed]
    keep_ids = all(i and re.fullmatch(r"SEC-\d+", i) for i in ids) and len(set(ids)) == len(ids)
    if "id" not in present or not keep_ids:
        notes.append("Security IDs were generated (SEC-001, SEC-002, …) so every fact can cite its row.")

    seen = {}
    for n, (xl_row, company, ticker, cell, vals) in enumerate(parsed, start=1):
        key = ticker.upper()
        if key in seen:
            raise UploadError([f"Ticker '{ticker}' appears more than once (rows {seen[key]} and {xl_row}). Tickers must be unique."], base_report)
        seen[key] = xl_row
        pts = {k: _pct(v / 100 if percent_points and v is not None else v) for k, v in vals.items() if k != "price"}
        price = vals["price"]
        complete = price is not None and all(pts[k] is not None for k in PERF_KEYS)
        score = None
        if complete:
            score = round(pts["d1"] * 0.05 + pts["m3"] * 0.2 + pts["m6"] * 0.3 + pts["y1"] * 0.45, 2)
        have = [price is not None] + [pts[k] is not None for k in PERF_KEYS]
        status = "complete" if complete else "missing" if not any(have) else "partial"
        url = _text(cell("source_url"), 300)
        rows.append({
            "id": ids[n - 1] if keep_ids else f"SEC-{n:03d}", "company": company, "ticker": ticker,
            "index": _text(cell("index"), 40), "sector": _text(cell("sector"), 60) or "Unassigned", "canton": _text(cell("canton"), 40),
            "isin": _text(cell("isin"), 20), "price": price, "d1": pts["d1"], "m3": pts["m3"], "m6": pts["m6"], "y1": pts["y1"],
            "score": score, "signal": signal_for(score), "signal_workbook": _text(cell("signal_workbook"), 20),
            "snapshot": _text(cell("snapshot"), 20), "coverage": _text(cell("coverage"), 60),
            "source_url": url if url and url.lower().startswith(("http://", "https://")) else None,
            "complete": complete, "status": status, "xl_row": xl_row,
        })
    names = defaultdict(list)
    for r in rows:
        names[r["company"]].append(r["ticker"])
    for r in rows:
        r["share_classes"] = names[r["company"]]

    if not rows:
        raise UploadError(["The sheet has a header row but no company rows."], base_report)
    if not any(r["complete"] for r in rows):
        raise UploadError(["No company has a price and all four performance values, so nothing can be analysed."], base_report)

    dates = Counter(d for d in (_parse_date(cell("snapshot")) for _r, _c, _t, cell, _v in parsed) if d)
    snapshot = dates.most_common(1)[0][0] if dates else date.today()
    if not dates:
        notes.append("No usable date column was found, so today's date is used as the snapshot date.")

    base_report["warnings"], base_report["notes"] = warnings, notes
    base_report["companies"] = [{k: r[k] for k in ("id", "company", "ticker", "sector", "price", "status", "share_classes", "xl_row")} for r in rows]
    base_report["counts"] = {
        "rows_read": len(records), "companies": len(rows), "with_data": sum(r["complete"] for r in rows),
        "no_data": sum(r["status"] != "complete" for r in rows), "distinct_companies": len(names),
        "share_class_groups": sum(1 for v in names.values() if len(v) > 1), "sectors": len({r["sector"] for r in rows}),
    }
    return Dataset(name, sha, len(data), is_sample, rows, snapshot, base_report)


def _mapping_report(headers, matches, records):
    out = []
    for i, h in enumerate(headers):
        if h is None or not str(h).strip():
            continue
        values = [r[i] for _n, r in records if i < len(r) and r[i] is not None and str(r[i]).strip()]
        numeric = sum(1 for v in values if isinstance(v, (int, float)) and not isinstance(v, bool))
        kind = "number" if values and numeric >= len(values) * 0.8 else "text"
        entry = {"column": str(h), "letter": openpyxl.utils.get_column_letter(i + 1), "samples": [_show(v) for v in values[:3]],
                 "filled": len(values), "blank": len(records) - len(values), "kind": kind}
        if i in matches:
            key, how = matches[i]
            entry |= {"field": key, "label": FIELD[key][1], "required": FIELD[key][2], "how": how}
        elif _norm(h) in RECOMPUTED_HEADERS:
            entry |= {"field": None, "label": "Momentum score", "required": False, "how": "recomputed"}
        else:
            entry |= {"field": None, "label": None, "required": False, "how": "ignored"}
        out.append(entry)
    return out


def _field_report(headers, matches):
    mapped = {k: headers[i] for i, (k, _how) in matches.items()}
    return [{"key": k, "label": label, "required": req, "desc": desc, "column": str(mapped[k]) if k in mapped else None}
            for k, label, req, _aliases, desc in FIELDS]


def _preview(headers, records):
    cols = [i for i, h in enumerate(headers) if h is not None and str(h).strip()][:9]
    return {"columns": [str(headers[i]) for i in cols],
            "rows": [[_show(r[i]) if i < len(r) else "" for i in cols] for _n, r in records[:6]]}


# ----------------------------------------------------------------------------- active dataset

_active = None
_override = ContextVar("dataset_override", default=None)


@lru_cache(maxsize=1)
def sample():
    return parse_workbook(XLSX_PATH.read_bytes(), SAMPLE_NAME, is_sample=True)


def current():
    return _override.get() or _active or sample()


def activate(ds):
    global _active
    _active = None if ds.is_sample else ds


def pin(ds):
    """Fixes the dataset for the current async task, so a running stream keeps its workbook when another file is uploaded."""
    _override.set(ds)


def pin_sample():
    """Makes the sample the dataset for the current async task only (Trust Lab and Model Bake-off are built on it)."""
    pin(sample())


# ----------------------------------------------------------------------------- analysis

def rows():
    return current().rows


def by_ticker(ticker):
    return current().index.get(str(ticker).upper())


def _avg(vals):
    vals = [v for v in vals if v is not None]
    return round(mean(vals), 2) if vals else None


def sector_stats():
    groups = defaultdict(list)
    for r in rows():
        groups[r["sector"]].append(r)
    out = []
    for sector, items in groups.items():
        done = [i for i in items if i["complete"]]
        out.append({
            "sector": sector, "n": len(items), "n_complete": len(done),
            "avg_d1": _avg([i["d1"] for i in done]), "avg_m3": _avg([i["m3"] for i in done]),
            "avg_m6": _avg([i["m6"] for i in done]), "avg_y1": _avg([i["y1"] for i in done]),
            "avg_score": _avg([i["score"] for i in done]),
            "positive": sum(1 for i in done if i["signal"] == "Positive"),
            "negative": sum(1 for i in done if i["signal"] == "Negative"),
        })
    return sorted(out, key=lambda s: -(s["avg_score"] if s["avg_score"] is not None else -999))


def sector_context(sector, ticker=None):
    peers = sorted((r for r in rows() if r["sector"] == sector and r["complete"]), key=lambda r: -r["score"])
    stats = next((s for s in sector_stats() if s["sector"] == sector), {})
    rank = None
    if ticker:
        rank = next((i + 1 for i, p in enumerate(peers) if p["ticker"] == ticker.upper()), None)
    return {
        "sector": sector, "stats": stats, "rank_in_sector": rank, "ranked_peers": len(peers),
        "top_peers": [{"ticker": p["ticker"], "company": p["company"], "score": p["score"]} for p in peers[:3]],
    }


def findings(today=None):
    today = today or date.today()
    ds = current()
    data = ds.rows
    complete = [r for r in data if r["complete"]]
    missing = [r for r in data if not r["complete"]]
    wb_known = [r for r in complete if r["signal_workbook"] in ("Positive", "Mixed", "Negative")]
    wb_pos = sum(1 for r in wb_known if r["signal_workbook"] == "Positive")
    wb_mixed = sum(1 for r in wb_known if r["signal_workbook"] == "Mixed")
    wb_neg = sum(1 for r in wb_known if r["signal_workbook"] == "Negative")
    pos = sum(1 for r in complete if r["signal"] == "Positive")
    neg = sum(1 for r in complete if r["signal"] == "Negative")
    mixed = len(complete) - pos - neg
    dup = defaultdict(list)
    for r in data:
        dup[r["company"]].append(r["ticker"])
    dups = {k: v for k, v in dup.items() if len(v) > 1}
    age = max(0, (today - ds.snapshot).days)
    out = []
    if wb_known and wb_pos + wb_neg == 0 and pos + neg > 0:
        out.append({
            "severity": "critical", "key": "scale", "title": "Signal formula is mis-scaled",
            "detail": f"The workbook shows {wb_mixed} of {len(complete)} securities as 'Mixed' and {wb_pos} as 'Positive'. "
                      f"Momentum Score multiplies decimal returns (e.g. 0.0636) but the thresholds (>= 20, <= -10) are in percentage points.",
            "action": f"Rescaled to percentage points: {pos} Positive, {mixed} Mixed, {neg} Negative.",
            "counts": {"positive": pos, "mixed": mixed, "negative": neg, "wb_mixed": wb_mixed, "wb_positive": wb_pos},
        })
    if missing:
        out.append({
            "severity": "warning", "key": "missing", "title": f"{len(missing)} {'security has' if len(missing) == 1 else 'securities have'} no complete market data",
            "detail": ", ".join(f"{r['company']} ({r['ticker']})" for r in missing[:8]) + (f" and {len(missing) - 8} more" if len(missing) > 8 else "") + ".",
            "action": "Excluded from rankings and averages. Never treated as 0%.",
            "tickers": [r["ticker"] for r in missing],
        })
    if dups:
        out.append({
            "severity": "info", "key": "share_classes", "title": f"{len(dups)} {'company appears' if len(dups) == 1 else 'companies appear'} twice (share classes)",
            "detail": "; ".join(f"{k}: {', '.join(v)}" for k, v in list(dups.items())[:6]) + ".",
            "action": f"Rankings group by company: {len(dup)} distinct companies instead of {len(data)} rows.",
        })
    out.append({
        "severity": "warning" if age > 14 else "info", "key": "stale", "title": f"Market snapshot is {age} days old",
        "detail": f"Prices and performance are as of {ds.snapshot_str}. Only price and performance columns are used; no revenue or earnings data.",
        "action": "Briefings must state the snapshot date and avoid claims about current performance.",
    })
    return out


def summary():
    ds = current()
    data = ds.rows
    complete = [r for r in data if r["complete"]]
    return {
        "snapshot": ds.snapshot_str,
        "source": ds.info(),
        "kpis": {
            "securities": len(data), "complete": len(complete),
            "distinct_companies": len({r["company"] for r in data}),
            "avg_m3": _avg([r["m3"] for r in complete]), "avg_m6": _avg([r["m6"] for r in complete]),
            "avg_y1": _avg([r["y1"] for r in complete]),
            "positive": sum(1 for r in complete if r["signal"] == "Positive"),
            "negative": sum(1 for r in complete if r["signal"] == "Negative"),
            "mixed": sum(1 for r in complete if r["signal"] == "Mixed"),
            "wb_positive": sum(1 for r in complete if r["signal_workbook"] == "Positive"),
            "wb_mixed": sum(1 for r in complete if r["signal_workbook"] == "Mixed"),
            "wb_negative": sum(1 for r in complete if r["signal_workbook"] == "Negative"),
        },
        "sectors": sector_stats(),
        "findings": findings(),
        "rows": data,
    }


@lru_cache(maxsize=1)
def demo_prompts():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["Demo Prompts"]
    keys = ["stage", "tool", "purpose", "prompt"]
    return [dict(zip(keys, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]


def llm_context_table(corrected=True):
    """Compact text table of the workbook used as grounding context for assistants."""
    lines = ["id | ticker | company | sector | price_chf | d1_% | m3_% | m6_% | y1_% | momentum_pts | signal"]
    for r in rows():
        def f(v):
            return "n/a" if v is None else str(v)
        sig = r["signal"] if corrected else (r["signal_workbook"] or "n/a")
        score = r["score"] if corrected else (None if r["score"] is None else round(r["score"] / 100, 4))
        lines.append(" | ".join([r["id"], r["ticker"], r["company"], r["sector"], f(r["price"]), f(r["d1"]),
                                 f(r["m3"]), f(r["m6"]), f(r["y1"]), f(score), sig]))
    return "\n".join(lines)
